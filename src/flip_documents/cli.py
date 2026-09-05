"""Original document capture and explicit local conversion, independent of Flip."""

from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
import re
import sys
import tempfile
from datetime import UTC, datetime
from email.message import Message
from pathlib import Path
from urllib.parse import unquote, urlsplit
from urllib.request import Request, urlopen

VERSION = "0.2.0"
USER_AGENT = f"flip-documents/{VERSION}"
TYPES = {
    "application/pdf": ".pdf",
    "text/csv": ".csv",
    "application/csv": ".csv",
    "application/zip": ".zip",
    "application/msword": ".doc",
    "application/vnd.ms-excel": ".xls",
    "application/vnd.ms-powerpoint": ".ppt",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
}
SUFFIXES = set(TYPES.values()) | {".odt", ".ods", ".xlsm"}


def filename(url: str, content_type: str, disposition: str) -> str:
    message = Message()
    message["Content-Disposition"] = disposition
    supplied = message.get_filename() or unquote(urlsplit(url).path).rsplit("/", 1)[-1]
    name = re.sub(
        r"[^A-Za-z0-9._-]", "_", supplied.replace("\\", "/").rsplit("/", 1)[-1]
    )
    name = name.lstrip(".")[:180] or "document"
    if name in {"flip.json", "receipt.json"}:
        name = "document"
    if Path(name).suffix.lower() not in SUFFIXES:
        name += TYPES.get(content_type, ".bin")
    return name


def capture(
    url: str, dest: Path, *, max_bytes: int = 100_000_000, timeout: float = 60
) -> dict:
    parsed = urlsplit(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username:
        raise ValueError("capture requires an HTTP(S) URL without embedded credentials")
    if max_bytes <= 0 or timeout <= 0:
        raise ValueError("limits must be positive")
    dest = dest.absolute()
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.exists() and (not dest.is_dir() or any(dest.iterdir())):
        raise ValueError("destination must be absent or an empty directory")
    with tempfile.TemporaryDirectory(
        prefix=".documents-", dir=dest.parent
    ) as temporary:
        staging = Path(temporary)
        request = Request(
            url, headers={"User-Agent": USER_AGENT, "Accept-Encoding": "identity"}
        )
        with urlopen(request, timeout=timeout) as response:
            final_url = response.geturl()
            if urlsplit(final_url).scheme not in {"http", "https"}:
                raise ValueError("redirect left HTTP(S)")
            content_type = response.headers.get_content_type()
            name = filename(
                final_url, content_type, response.headers.get("Content-Disposition", "")
            )
            first = response.read(min(8192, max_bytes + 1))
            head = first.lstrip().lower()
            if (
                not first
                or head.startswith((b"<!doctype html", b"<html", b"<?xml"))
                or content_type in {"text/html", "application/xhtml+xml"}
            ):
                raise ValueError(
                    "response is empty or a page, not the requested document"
                )
            magic = first.startswith((b"%PDF-", b"PK\x03\x04", b"\xd0\xcf\x11\xe0"))
            suffix = Path(unquote(urlsplit(final_url).path)).suffix.lower()
            if (
                content_type not in TYPES
                and not magic
                and not (
                    content_type == "application/octet-stream" and suffix in SUFFIXES
                )
            ):
                raise ValueError(f"unrecognized document content type: {content_type}")
            if first.startswith(b"%PDF-"):
                name = str(Path(name).with_suffix(".pdf"))
                content_type = "application/pdf"
            digest = hashlib.sha256()
            size = 0
            with (staging / name).open("wb") as output:
                chunk = first
                while chunk:
                    size += len(chunk)
                    if size > max_bytes:
                        raise ValueError(f"document exceeds {max_bytes} byte limit")
                    digest.update(chunk)
                    output.write(chunk)
                    chunk = response.read(min(65536, max_bytes - size + 1))
            declared = response.headers.get("Content-Length")
            if declared is not None and int(declared) != size:
                raise ValueError(
                    "incomplete response: Content-Length does not match received bytes"
                )
            receipt = {
                "schema": "flip.documents-capture/1",
                "input_url": url,
                "canonical_url": final_url,
                "retrieved_at": datetime.now(UTC).isoformat(),
                "filename": name,
                "sha256": digest.hexdigest(),
                "bytes": size,
                "mime": content_type,
                "strategy": "http-get",
                "status": "success",
                "user_agent": USER_AGENT,
                "tool": USER_AGENT,
                "content_encoding": response.headers.get(
                    "Content-Encoding", "identity"
                ),
            }
            if receipt["content_encoding"].lower() not in {"identity", ""}:
                raise ValueError(
                    "server ignored identity encoding; refusing opaque encoded document"
                )
        (staging / "flip.json").write_text(
            json.dumps({"flip": receipt}, indent=2) + "\n", encoding="utf-8"
        )
        if dest.exists():
            dest.rmdir()
        staging.rename(dest)
    return receipt


def extract(src: Path, out: Path) -> dict:
    if src.resolve() == out.resolve() or out.exists():
        raise ValueError("output must be new and distinct from the original")
    if src.suffix.lower() not in {".docx", ".xlsx", ".pptx", ".pdf"}:
        raise ValueError("conversion supports DOCX, XLSX, PPTX and text-layer PDF only")
    try:
        from markitdown import MarkItDown
    except ImportError as exc:
        raise ValueError(
            'install extraction support: pip install "flip-documents[markitdown]"'
        ) from exc
    result = MarkItDown(enable_plugins=False).convert(str(src))
    text = result.markdown
    if not text.strip():
        raise ValueError("converter found no readable text; no derivative written")
    out.parent.mkdir(parents=True, exist_ok=True)
    # Exclusive creation preserves existing work even if it appeared during conversion.
    with out.open("x", encoding="utf-8") as output:
        output.write(text + "\n")
    return {
        "schema": "flip.documents-extraction/1",
        "tool": "markitdown",
        "version": importlib.metadata.version("markitdown"),
        "method": "text-layer",
        "input_sha256": hashlib.sha256(src.read_bytes()).hexdigest(),
        "output_sha256": hashlib.sha256(out.read_bytes()).hexdigest(),
        "coverage": "reading derivative; no OCR, formula evaluation, or completeness guarantee",
    }


def cells(src: Path, out: Path) -> dict:
    """Read workbook cells through openpyxl without pandas type inference."""
    if src.resolve() == out.resolve() or out.exists():
        raise ValueError("output must be new and distinct from the original")
    if src.suffix.lower() != ".xlsx":
        raise ValueError("cell extraction supports XLSX only")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError(
            'install spreadsheet support: pip install "flip-documents[markitdown]"'
        ) from exc
    source_hash = hashlib.sha256(src.read_bytes()).hexdigest()
    workbook = load_workbook(src, read_only=True, data_only=False, keep_links=False)
    cached = None
    sheets = []
    formulas = 0
    try:
        cached = load_workbook(src, read_only=True, data_only=True, keep_links=False)
        for sheet in workbook:
            if sheet.max_row is None or sheet.max_column is None:
                raise ValueError(
                    "worksheet has no declared dimensions; inspect it in a spreadsheet tool"
                )
            if sheet.max_row * sheet.max_column > 1_000_000:
                raise ValueError("worksheet declared range exceeds 1,000,000 cells")
            records = []
            for row, cached_row in zip(
                sheet.iter_rows(), cached[sheet.title].iter_rows(), strict=True
            ):
                for cell, cached_cell in zip(row, cached_row, strict=True):
                    if cell.value is None:
                        continue
                    record = {
                        "coordinate": cell.coordinate,
                        "data_type": cell.data_type,
                        "value": cell.value,
                        "number_format": cell.number_format,
                    }
                    if cell.data_type == "f":
                        formulas += 1
                        record["cached_value"] = cached_cell.value
                        record["cache_status"] = (
                            "present"
                            if cached_cell.value is not None
                            else "unavailable-or-blank"
                        )
                    records.append(record)
            sheets.append(
                {"name": sheet.title, "state": sheet.sheet_state, "cells": records}
            )
    finally:
        workbook.close()
        if cached is not None:
            cached.close()
    if hashlib.sha256(src.read_bytes()).hexdigest() != source_hash:
        raise ValueError("original changed during extraction; no derivative written")
    backend_version = importlib.metadata.version("openpyxl")
    artifact = {
        "schema": "flip.documents-cells/1",
        "input_sha256": source_hash,
        "tool": "openpyxl",
        "version": backend_version,
        "sheets": sheets,
        "limitations": [
            "Formulas are not evaluated; cached values may be stale or unavailable.",
            "Dates and times use ISO text; number_format retains display instructions.",
            "Numbers are decoded Python integers/floats, not exact XML decimal text.",
            "Blank cells are omitted; coordinates preserve gaps. Styles, merged ranges, charts and comments are not exported.",
        ],
    }

    # openpyxl decodes Excel date/time values; retain their ISO representation.
    def encode(value):
        if hasattr(value, "isoformat"):
            return value.isoformat()
        raise TypeError(f"unsupported cell value type: {type(value).__name__}")

    text = json.dumps(artifact, indent=2, ensure_ascii=False, default=encode) + "\n"
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("x", encoding="utf-8") as output:
        output.write(text)
    return {
        "schema": "flip.documents-extraction/1",
        "tool": "openpyxl",
        "version": backend_version,
        "method": "structured",
        "input_sha256": source_hash,
        "output_sha256": hashlib.sha256(out.read_bytes()).hexdigest(),
        "formulas": formulas,
    }


def version() -> str:
    try:
        backend = importlib.metadata.version("markitdown")
    except importlib.metadata.PackageNotFoundError:
        backend = "not-installed"
    try:
        spreadsheet = importlib.metadata.version("openpyxl")
    except importlib.metadata.PackageNotFoundError:
        spreadsheet = "not-installed"
    return f"flip-documents {VERSION}; markitdown {backend}; openpyxl {spreadsheet}"


def plugin_path() -> Path:
    bundled = Path(__file__).resolve().parent / "bundle"
    return bundled if bundled.is_dir() else Path(__file__).resolve().parents[2]


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--version", action="version", version=version())
    parser.add_argument(
        "--plugin-path",
        action="store_true",
        help="print the bundled Flip plugin directory",
    )
    commands = parser.add_subparsers(dest="command")
    fetch = commands.add_parser(
        "capture", help="save original HTTP document and receipt"
    )
    fetch.add_argument("url")
    fetch.add_argument("dest", type=Path)
    fetch.add_argument("--max-bytes", type=int, default=100_000_000)
    fetch.add_argument("--timeout", type=float, default=60)
    derive = commands.add_parser(
        "extract", help="convert a local original using optional MarkItDown"
    )
    derive.add_argument("src", type=Path)
    derive.add_argument("out", type=Path)
    spreadsheet = commands.add_parser(
        "cells", help="preserve XLSX coordinates, types, formulas and caches as JSON"
    )
    spreadsheet.add_argument("src", type=Path)
    spreadsheet.add_argument("out", type=Path)
    args = parser.parse_args(argv)
    if args.plugin_path:
        print(plugin_path())
        return 0
    if args.command is None:
        parser.error("choose capture, extract or cells")
    try:
        result = (
            capture(args.url, args.dest, max_bytes=args.max_bytes, timeout=args.timeout)
            if args.command == "capture"
            else (
                cells(args.src, args.out)
                if args.command == "cells"
                else extract(args.src, args.out)
            )
        )
        print(json.dumps(result, ensure_ascii=False))
        return 0
    except Exception as exc:  # noqa: BLE001 - CLI boundary includes optional converter errors
        print(f"flip-documents: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
