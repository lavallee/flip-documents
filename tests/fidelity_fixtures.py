"""Original synthetic fixtures; no publisher documents or third-party assets."""

from pathlib import Path
from zipfile import ZipFile


def docx(path: Path):
    with ZipFile(path, "w") as z:
        z.writestr(
            "[Content_Types].xml",
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
        )
        z.writestr(
            "_rels/.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>',
        )
        z.writestr(
            "word/document.xml",
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>First paragraph: café &amp; library.</w:t></w:r></w:p><w:tbl><w:tr><w:tc><w:p><w:r><w:t>Code</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>00123</w:t></w:r></w:p></w:tc></w:tr></w:tbl><w:p><w:r><w:t>Last paragraph: negative -12.50.</w:t></w:r></w:p></w:body></w:document>',
        )


def xlsx(path: Path):
    from openpyxl import Workbook

    w = Workbook()
    s = w.active
    s.title = "Sparse"
    s.append(["Code", "Amount", "Gap", "Total"])
    s["A2"] = "00123"
    s["B2"] = 12.5
    s["B2"].number_format = "0.00"
    s["D2"] = "=B2*2"
    s["E1"] = "Cached formula"
    s["E2"] = "=B2*2"
    s["A4"] = "00999"
    s["B4"] = -1.25
    s["D4"] = 0
    w.create_sheet("Second")["A1"] = "SECOND SHEET"
    w.save(path)
    # Synthetic publisher-style cache: openpyxl deliberately does not calculate it.
    from xml.etree import ElementTree as ET

    with ZipFile(path) as source:
        members = {name: source.read(name) for name in source.namelist()}
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    tree = ET.fromstring(members["xl/worksheets/sheet1.xml"])
    tree.find(f'.//{ns}c[@r="E2"]/{ns}v').text = "25"
    members["xl/worksheets/sheet1.xml"] = ET.tostring(tree)
    with ZipFile(path, "w") as target:
        for name, content in members.items():
            target.writestr(name, content)


def pptx(path: Path):
    from pptx import Presentation
    from pptx.util import Inches

    p = Presentation()
    for i in range(1, 13):
        slide = p.slides.add_slide(p.slide_layouts[6])
        slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(5), Inches(1)
        ).text = f"Slide original {i}"
    # Presentation order deliberately differs from slide XML filename order.
    last = p.slides._sldIdLst[-1]
    p.slides._sldIdLst.remove(last)
    p.slides._sldIdLst.insert(0, last)
    p.save(path)


def pdf(path: Path):
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R 5 0 R] /Count 2 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 7 0 R >> >> /Contents 4 0 R >>",
        None,
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 7 0 R >> >> /Contents 6 0 R >>",
        None,
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for position, text in [
        (3, "Page ONE: code 00123 and amount -12.50."),
        (5, "Page TWO: library open on Tuesday."),
    ]:
        stream = f"BT /F1 12 Tf 50 730 Td ({text}) Tj ET".encode()
        objects[position] = (
            f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream"
        )
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(data)
    data.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    path.write_bytes(data)


def scan(path: Path):
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (600, 800), "white")
    ImageDraw.Draw(image).text((50, 50), "SCAN ONLY: code 00123", fill="black")
    image.save(path, "PDF")
